import hashlib
import io
import json
import random
import re
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openai import (
    APIConnectionError,
    APIStatusError,
    APITimeoutError,
    AuthenticationError,
    OpenAI,
    RateLimitError,
)

APP_TITLE = "NurseInsight - 护理不良事件智能RCA辅助系统"
DISCLAIMER = "本系统分析结果仅为护士长/质控护士提供辅助参考，最终决策权和解释权归临床专业人员所有。"
REPORT_TITLE = "护理不良事件分析报告"

DATA_PAGE = "📁 数据管理与录入"
ANALYSIS_PAGE = "🧠 智能根因分析"
DASHBOARD_PAGE = "📊 可视化质控看板"

PAGES = [DATA_PAGE, ANALYSIS_PAGE, DASHBOARD_PAGE]

REQUIRED_COLUMNS = ["事件ID", "发生日期", "科室", "事件类型", "事件经过描述"]

DEPARTMENTS = [
    "心血管内科",
    "神经内科",
    "呼吸内科",
    "消化内科",
    "普通外科",
    "骨科",
    "重症医学科",
    "急诊科",
    "儿科",
    "产科",
    "肿瘤科",
    "肾内科",
]

SEVERITY_LEVELS = ["轻微", "隐患", "警告", "中度", "重度"]

RCA_DIMENSIONS = [
    "人（人为因素）",
    "机（设备设施）",
    "料（药品耗材）",
    "法（制度流程）",
    "环（环境管理）",
]

EVENT_TEMPLATES = {
    "跌倒/坠床": [
        "凌晨2时15分患者自行下床如厕，未按呼叫器求助，在卫生间门口滑倒，左侧髋部着地，立即予卧床制动并通知值班医师，X线检查未见骨折",
        "患者趁家属外出打水期间翻越床档下床，双膝跪地跌倒，左膝皮肤擦伤约2cm，予碘伏消毒包扎，加强宣教并落实床档使用",
        "术后第2天首次下床活动，因体位性低血压出现头晕，护士扶助不及时致患者坐倒在地，血压92/58mmHg，予平卧休息30分钟后缓解",
        "夜班巡视间隔期间，患者从无床档的病床坠床，头部撞击地面，即刻查体神志清楚，行头颅CT未见异常，已上报护理部",
    ],
    "给药错误": [
        "夜班护士在给3床患者发药时，因同时接听电话导致发错药，患者服用后出现轻微不良反应",
        "护士将12床患者的胰岛素误发予同姓的21床患者，患者于餐前皮下注射8U，30分钟后测血糖2.9mmol/L，立即口服葡萄糖并持续监测",
        "责任护士摆药时未核对床号，将鼻饲用药误经静脉通路注入，发现后立即停止输注，更换输液器并报告医师",
        "护士未严格执行查对制度，将缓释片研磨后给患者鼻饲，导致药物血药浓度波动，医师重新调整用药方案",
        "实习护士在带教老师不在场的情况下为患者静推药物，剂量计算错误，患者出现心慌、出汗，予对症处理后缓解",
    ],
    "管路滑脱": [
        "患者夜间烦躁不安，在未使用约束带的情况下自行拔除胃管，予重新置管并加强宣教及家属陪护",
        "护士为患者翻身时未妥善固定引流管，导致腹腔引流管部分脱出约5cm，立即通知医师评估，予重新固定并观察引流液性状",
        "患者因剧烈咳嗽致气管插管移位，SpO2下降至88%，立即予吸痰、调整导管深度，复测SpO2回升至96%",
        "清醒患者自行拔除PICC导管，穿刺点少量渗血，予压迫止血、更换敷料，评估血管条件后重新置管",
    ],
    "压力性损伤": [
        "患者因脑梗死后长期卧床，交接班时发现骶尾部皮肤发红、压之不褪色，属1期压力性损伤，予气垫床减压并每2小时翻身",
        "ICU患者使用无创呼吸机面罩过紧，鼻梁处出现2期压力性损伤，予泡沫敷料保护并调整面罩固定带松紧度",
    ],
    "标本采集错误": [
        "护士未核对检验条码，将两名患者的血标本送检顺序颠倒，检验科电话反馈后立即重新采集并更正报告",
        "患者输血前血型鉴定标本标签粘贴错误，被检验科拒收后重新采血，导致输血时间延迟1小时",
    ],
    "输液外渗": [
        "患者输注去甲肾上腺素过程中穿刺处出现肿胀、皮肤苍白，立即停止输注并拔针，予酚妥拉明局部封闭",
        "老年患者外周静脉条件差，输注甘露醇时发生外渗，局部组织肿胀约3×4cm，予硫酸镁湿敷并抬高患肢",
    ],
}

SEVERITY_WEIGHTS = {
    "跌倒/坠床": [0.15, 0.25, 0.20, 0.28, 0.12],
    "给药错误": [0.30, 0.28, 0.24, 0.13, 0.05],
    "管路滑脱": [0.18, 0.22, 0.30, 0.22, 0.08],
    "压力性损伤": [0.35, 0.30, 0.20, 0.15, 0.00],
    "标本采集错误": [0.40, 0.30, 0.20, 0.10, 0.00],
    "输液外渗": [0.32, 0.26, 0.24, 0.18, 0.00],
}

PROJECT_DIR = Path(__file__).resolve().parent
# 打包分享前必须删除 config.json，或改名为 config.example.json 并清空密钥
CONFIG_PATH = PROJECT_DIR / "config.json"
FEEDBACK_PATH = PROJECT_DIR / "user_feedback.json"
API_KEY_FIELD = "DEEPSEEK_API_KEY"
API_KEY_PLACEHOLDER = "我的密钥"

EDITOR_KEY_PREFIX = "edit_"

DEEPSEEK_BASE_URL = "https://api.deepseek.com"
DEEPSEEK_MODEL = "deepseek-chat"
API_TIMEOUT_SECONDS = 60

SEVERITY_BADGES = {
    "轻微": "🟢 轻微",
    "隐患": "🟡 隐患",
    "警告": "🟠 警告",
    "中度": "🔴 中度",
    "重度": "🟣 重度",
}

OTHER_LABEL = "其他"
MAX_PIE_SLICES = 6
CHART_HEIGHT = 420
FONT_FAMILY = 'system-ui, -apple-system, "Segoe UI", sans-serif'

DEPARTMENT_CHART_TITLE = "各科室不良事件数量"
TYPE_CHART_TITLE = "不良事件类型占比"
SEVERITY_CHART_TITLE = "严重程度分布"
FALLBACK_TYPE_CHART_TITLE = "事件类型分布（数据无「严重程度」字段）"

# 序数色阶按严重程度递增提高对比度：深色画布上越亮，浅色画布上越深
SCREEN_STYLE = {
    "surface": "#141B24",
    "ink": "#F2F6FA",
    "muted": "#8A98A8",
    "grid": "#2A3440",
    "axis": "#3A4654",
    "palette": ["#3987e5", "#d95926", "#199e70", "#c98500", "#d55181", "#008300"],
    "ordinal": ["#184f95", "#256abf", "#3987e5", "#6da7ec", "#9ec5f4"],
    "other": "#8A98A8",
}

DOCUMENT_STYLE = {
    "surface": "#FFFFFF",
    "ink": "#0B0B0B",
    "muted": "#898781",
    "grid": "#E1E0D9",
    "axis": "#C3C2B7",
    "palette": ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300"],
    "ordinal": ["#86b6ef", "#5598e7", "#2a78d6", "#1c5cab", "#104281"],
    "other": "#898781",
}

DOC_ACCENT_COLOR = RGBColor(0x1C, 0x5C, 0xAB)
DOC_MUTED_COLOR = RGBColor(0x52, 0x51, 0x4E)
DOC_BODY_FONT = "宋体"
DOC_HEADING_FONT = "微软雅黑"
DOC_LATIN_FONT = "Calibri"
DOC_IMAGE_WIDTH_INCHES = 6.0
CHART_IMAGE_WIDTH_PX = 1000
CHART_IMAGE_HEIGHT_PX = 600
CHART_IMAGE_SCALE = 2

MISSING_THEME_HINT = (
    "⚠️ 未检测到深色主题配置：请保留项目内的 .streamlit/config.toml，"
    "否则表格与输入框会回退成浅色，出现浅底浅字的显示问题。"
)

GLOBAL_CSS = """
<style>
:root {
  color-scheme: dark;
  --ni-bg: #0E1117;
  --ni-surface: #161D26;
  --ni-surface-2: #141B24;
  --ni-input: #0F151D;
  --ni-border: #26313D;
  --ni-border-strong: #33455A;
  --ni-accent: #3987e5;
  --ni-accent-soft: rgba(57, 135, 229, 0.14);
  --ni-ink: #E6EDF5;
  --ni-ink-muted: #9FB0C3;
  --ni-shadow: 0 0 0 1px rgba(57, 135, 229, 0.10), 0 8px 24px rgba(0, 0, 0, 0.45), 0 0 18px rgba(57, 135, 229, 0.07);
}

html, body, .stApp, [data-testid="stAppViewContainer"] {
  background-color: var(--ni-bg) !important;
  color: var(--ni-ink) !important;
}
[data-testid="stHeader"] { background-color: rgba(14, 17, 23, 0.85) !important; }
[data-testid="stToolbar"] { right: 1rem; }

.stApp :is(p, li, label, strong, em, td, th, caption) { color: var(--ni-ink); }
h1 { color: #F4F8FD !important; letter-spacing: 0.4px; }
h2, h3, h4 { color: #CFE0F2 !important; }
[data-testid="stCaptionContainer"] :is(p, span) { color: var(--ni-ink-muted) !important; }
hr { border-color: var(--ni-border) !important; }
a { color: var(--ni-accent) !important; }
code { background-color: #1B2430 !important; color: #A8D0FF !important; }

[data-testid="stSidebar"] {
  background-color: #131A24 !important;
  border-right: 1px solid var(--ni-border);
}
[data-testid="stSidebar"] :is(p, li, span, label, h1, h2, h3) { color: var(--ni-ink); }
[data-testid="stSidebar"] h3 { color: var(--ni-accent) !important; }
[data-testid="stSidebar"] hr { margin: 0.6rem 0; border-color: var(--ni-border) !important; }

[data-testid="stSidebar"] .stRadio div[role="radiogroup"] { gap: 8px; }
[data-testid="stSidebar"] .stRadio div[role="radiogroup"] > label {
  display: flex;
  align-items: center;
  padding: 10px 12px;
  background-color: #18212C;
  border: 1px solid var(--ni-border);
  border-left: 3px solid var(--ni-border);
  border-radius: 10px;
  cursor: pointer;
  transition: background-color 0.18s ease, border-color 0.18s ease, transform 0.18s ease, box-shadow 0.18s ease;
}
[data-testid="stSidebar"] .stRadio div[role="radiogroup"] > label:hover {
  background-color: var(--ni-accent-soft);
  border-color: var(--ni-accent);
  border-left-color: var(--ni-accent);
  transform: translateX(3px);
  box-shadow: 0 0 0 1px rgba(57, 135, 229, 0.35), 0 0 14px rgba(57, 135, 229, 0.25);
}
[data-testid="stSidebar"] .stRadio div[role="radiogroup"] > label:has(input:checked),
[data-testid="stSidebar"] .stRadio div[role="radiogroup"] > label[aria-checked="true"] {
  background-color: var(--ni-accent-soft);
  border-color: var(--ni-accent);
  border-left-color: var(--ni-accent);
  box-shadow: inset 0 0 18px rgba(57, 135, 229, 0.12);
  font-weight: 600;
}
[data-testid="stSidebar"] .stRadio div[role="radiogroup"] > label:has(input:checked) p { color: #FFFFFF !important; }

.stTextArea textarea, .stTextInput input, [data-baseweb="textarea"], [data-baseweb="input"] {
  background-color: var(--ni-input) !important;
  color: var(--ni-ink) !important;
  border-color: var(--ni-border) !important;
}
.stTextArea textarea, .stTextInput input {
  border: 1px solid var(--ni-border) !important;
  border-radius: 10px !important;
  caret-color: var(--ni-accent);
}
.stTextArea textarea:focus, .stTextInput input:focus {
  border-color: var(--ni-accent) !important;
  box-shadow: 0 0 0 1px rgba(57, 135, 229, 0.45) !important;
}
.stTextArea textarea::placeholder, .stTextInput input::placeholder {
  color: #6F8196 !important;
  opacity: 1;
}
[data-baseweb="select"] > div {
  background-color: var(--ni-input) !important;
  border-color: var(--ni-border) !important;
  color: var(--ni-ink) !important;
}
[data-baseweb="select"] svg { fill: var(--ni-ink-muted) !important; }
[data-baseweb="tag"] {
  background-color: rgba(57, 135, 229, 0.22) !important;
  border: 1px solid rgba(57, 135, 229, 0.45) !important;
}
[data-baseweb="tag"] :is(span, div) { color: #D7E7FB !important; }
[data-baseweb="popover"] :is(ul, li, div) {
  background-color: var(--ni-surface) !important;
  color: var(--ni-ink) !important;
}
[data-baseweb="popover"] li:hover, [data-baseweb="popover"] [role="option"]:hover {
  background-color: var(--ni-accent-soft) !important;
}
[data-baseweb="calendar"] :is(div, span, button) { color: var(--ni-ink) !important; }
[data-baseweb="calendar"] { background-color: var(--ni-surface) !important; }

[data-testid="stDataFrame"], [data-testid="stDataFrameResizable"], [data-testid="stDataFrameContainer"] {
  background-color: var(--ni-surface-2) !important;
  border: 1px solid var(--ni-border) !important;
  border-radius: 12px !important;
  overflow: hidden;
}
[data-testid="stDataFrame"] :is(div, span) { color: var(--ni-ink) !important; }
[data-testid="stTable"], [data-testid="stTable"] table {
  background-color: var(--ni-surface-2) !important;
  color: var(--ni-ink) !important;
  border-radius: 12px;
}
[data-testid="stTable"] :is(th, td) {
  background-color: var(--ni-surface-2) !important;
  color: var(--ni-ink) !important;
  border-color: var(--ni-border) !important;
}

[data-testid="stMetric"] {
  background-color: var(--ni-surface-2) !important;
  border: 1px solid var(--ni-border-strong);
  border-radius: 14px;
  padding: 14px 16px;
  box-shadow: var(--ni-shadow);
}
[data-testid="stMetric"] :is([data-testid="stMetricLabel"], [data-testid="stMetricLabel"] *) {
  color: var(--ni-ink-muted) !important;
}
[data-testid="stMetricValue"], [data-testid="stMetricValue"] * { color: #EAF2FB !important; }
[data-testid="stMetricDelta"] { color: #6FD08C !important; }

[data-testid="stVerticalBlockBorderWrapper"] {
  background-color: var(--ni-surface-2) !important;
  border: 1px solid var(--ni-border-strong) !important;
  border-radius: 14px;
  box-shadow: var(--ni-shadow);
  padding: 6px 12px;
}

[data-testid="stExpander"] details {
  background-color: #131A24 !important;
  border: 1px solid var(--ni-border) !important;
  border-radius: 12px !important;
}
[data-testid="stExpander"] summary { color: var(--ni-ink) !important; }
[data-testid="stExpander"] summary:hover { color: var(--ni-accent) !important; }
[data-testid="stExpander"] :is(p, span, li, code) { color: var(--ni-ink) !important; }

[data-testid="stJson"] {
  background-color: var(--ni-input) !important;
  border: 1px solid var(--ni-border) !important;
  border-radius: 10px;
}

[data-testid="stBaseButton-primary"], button[kind="primary"] {
  border-radius: 10px !important;
  font-weight: 600 !important;
  transition: box-shadow 0.18s ease, transform 0.18s ease;
}
[data-testid="stBaseButton-primary"]:hover, button[kind="primary"]:hover {
  box-shadow: 0 0 0 1px rgba(57, 135, 229, 0.5), 0 4px 16px rgba(57, 135, 229, 0.35);
  transform: translateY(-1px);
}
[data-testid="stBaseButton-secondary"], [data-testid="stDownloadButton"] button {
  background-color: #18212C !important;
  color: var(--ni-ink) !important;
  border: 1px solid var(--ni-border-strong) !important;
  border-radius: 10px !important;
  transition: border-color 0.18s ease, color 0.18s ease, box-shadow 0.18s ease;
}
[data-testid="stBaseButton-secondary"]:hover, [data-testid="stDownloadButton"] button:hover {
  border-color: var(--ni-accent) !important;
  color: #CFE6FF !important;
  box-shadow: 0 0 12px rgba(57, 135, 229, 0.25);
}

[data-testid="stAlert"] {
  border-radius: 12px;
  border: 1px solid var(--ni-border);
}

[data-testid="stFileUploader"] section {
  background-color: var(--ni-surface-2) !important;
  border: 1px dashed var(--ni-border-strong) !important;
  border-radius: 12px;
}
[data-testid="stFileUploader"] :is(span, small, div) { color: var(--ni-ink-muted) !important; }

::-webkit-scrollbar { width: 10px; height: 10px; }
::-webkit-scrollbar-track { background: #0E1117; }
::-webkit-scrollbar-thumb { background: #2A3A4D; border-radius: 6px; }
::-webkit-scrollbar-thumb:hover { background: #3A4E63; }
</style>
"""

SYSTEM_PROMPT = """你是医院护理部的"护理质控辅助工具"，专门协助护士长与质控护士对护理不良事件进行根本原因分析（RCA）。

分析必须严格依据"人、机、料、法、环（4M1E）"框架：
- 人（人为因素）：当事人资质、经验、疲劳度、查对执行、沟通协作、培训掌握程度
- 机（设备设施）：仪器设备状态、报警设置、床档/约束具/输液泵等器械的可用性与维护
- 料（药品耗材）：药品外观相似度、标签标识、耗材规格匹配、备药与存放方式
- 法（制度流程）：核心制度落实（查对、交接班、分级护理、巡视）、流程设计缺陷、人力排班、带教监督
- 环（环境管理）：照明、地面、空间布局、噪音干扰、陪护管理、工作被打断

工作规则：
1. 先判断输入内容是否属于护理不良事件描述（临床场景、涉及患者安全、有具体经过或后果）。
2. 若输入与护理不良事件无关（例如闲聊、"今天天气真好"、与医疗无关的内容），必须返回 is_event 为 false，并在 message 中礼貌说明本系统仅用于护理不良事件分析，同时给出输入示例，绝不强行分析。
3. 若属于不良事件，必须深挖管理层面的系统性漏洞，不得把原因简单归咎于"护士责任心不强"。

只输出一个 JSON 对象，不要输出任何解释性文字或 Markdown 代码块，结构如下：
{
  "is_event": true,
  "message": "",
  "event_type": "事件类型分类，如 给药错误、跌倒/坠床、管路滑脱 等",
  "severity": "只能取 警告、隐患、轻微、中度、重度 之一",
  "root_causes": {
    "人（人为因素）": ["具体原因，每条不超过30字"],
    "机（设备设施）": ["具体原因"],
    "料（药品耗材）": ["具体原因"],
    "法（制度流程）": ["具体原因"],
    "环（环境管理）": ["具体原因"]
  },
  "deep_analysis": "一段话深度剖析事件背后的管理漏洞，150字以内",
  "recommendations": [
    {"措施": "可落地的改进措施", "优先级": "高/中/低"}
  ]
}

当 is_event 为 false 时，只填写 is_event 与 message，其余字段留空："event_type" 为空字符串，"severity" 为空字符串，"root_causes" 为空对象，"deep_analysis" 为空字符串，"recommendations" 为空数组。
recommendations 至少给出3条，优先级只能是 高、中、低 三者之一。"""

RECTIFICATION_PROMPT = """你是医院护理部的质控专家，负责为护理不良事件制定整改方案。

要求：
1. 输出 200 字以内的整改建议，分为 3 至 4 条，每条单独成段。
2. 每条直接给出可落地的动作（做什么、谁来落实、做到什么标准），不要空话套话。
3. 聚焦制度流程、培训考核、人力配置、设备环境等系统层面，避免只强调个人责任心。
4. 只输出建议正文，不要标题、不要 Markdown 代码块、不要序号前缀以外的修饰符号。"""


@st.cache_data
def generateEvents(count=100):
    rng = random.Random(20260101)
    startDay = date(2026, 1, 1)
    eventDays = sorted(
        startDay + timedelta(days=rng.randint(0, 272)) for _ in range(count)
    )

    rows = []
    for index, occurredOn in enumerate(eventDays):
        eventType = rng.choice(list(EVENT_TEMPLATES.keys()))
        rows.append(
            {
                "事件ID": f"AE-{occurredOn.year}-{index + 1:04d}",
                "发生日期": pd.to_datetime(occurredOn),
                "科室": rng.choice(DEPARTMENTS),
                "事件类型": eventType,
                "严重程度": rng.choices(SEVERITY_LEVELS, weights=SEVERITY_WEIGHTS[eventType])[0],
                "事件经过描述": rng.choice(EVENT_TEMPLATES[eventType]),
            }
        )

    return pd.DataFrame(rows)


def getEventData():
    if "eventData" in st.session_state:
        return st.session_state["eventData"]
    return generateEvents()


def loadConfiguredApiKey():
    if CONFIG_PATH.is_file():
        try:
            config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            config = {}
        apiKey = str(config.get(API_KEY_FIELD, "")).strip()
        if apiKey and apiKey != API_KEY_PLACEHOLDER:
            return apiKey, "config.json"

    try:
        secretKey = str(st.secrets[API_KEY_FIELD]).strip()
    except (KeyError, FileNotFoundError):
        return None, None

    if secretKey and secretKey != API_KEY_PLACEHOLDER:
        return secretKey, "Streamlit Secrets"
    return None, None


def readUploadedFile(uploadedFile):
    if not uploadedFile.name.lower().endswith((".xlsx", ".xls")):
        st.error(f"文件格式不支持：{uploadedFile.name}。请上传 .xlsx 或 .xls 格式的 Excel 表格。")
        return None

    try:
        return pd.read_excel(uploadedFile)
    except (ValueError, OSError, ImportError, zipfile.BadZipFile) as error:
        st.error(f"读取失败，文件可能已损坏、被加密或不是真正的 Excel 文件。详细信息：{error}")
        return None


def formatDataSourceLabel():
    if "dataSource" in st.session_state:
        return f"上传文件 {st.session_state['dataSource']}"
    return "系统生成的模拟数据（100 条）"


def renderDataPage():
    st.header("数据管理与录入")
    st.info("本模块用于导入护理不良事件原始记录（Excel），提供字段映射、缺失值校验与数据清洗。")

    uploadedFile = st.file_uploader("上传护理不良事件 Excel 表格", type=["xlsx", "xls"])

    if uploadedFile is not None:
        uploadedFrame = readUploadedFile(uploadedFile)
        if uploadedFrame is not None:
            if uploadedFrame.empty:
                st.warning("上传的表格没有任何数据行，已继续使用模拟数据。")
            else:
                missingColumns = [c for c in REQUIRED_COLUMNS if c not in uploadedFrame.columns]
                if missingColumns:
                    st.warning(
                        "上传的表格缺少必要字段："
                        + "、".join(missingColumns)
                        + "。已继续使用模拟数据。"
                    )
                else:
                    st.session_state["eventData"] = uploadedFrame
                    st.session_state["dataSource"] = uploadedFile.name
                    st.success(f"已成功载入 {len(uploadedFrame)} 条记录。")

    if "eventData" in st.session_state:
        eventFrame = getEventData()
        st.caption(f"当前数据来源：{formatDataSourceLabel()}")
        if st.button("恢复模拟数据"):
            st.session_state.pop("eventData", None)
            st.session_state.pop("dataSource", None)
            st.rerun()
    else:
        eventFrame = getEventData()
        st.caption("当前数据来源：系统生成的模拟数据（100 条）")

    st.subheader(f"事件记录（共 {len(eventFrame)} 条）")
    st.dataframe(
        eventFrame,
        width="stretch",
        hide_index=True,
        column_config={
            "事件经过描述": st.column_config.TextColumn(width="large"),
            "发生日期": st.column_config.DateColumn(format="YYYY-MM-DD"),
        },
    )


def parseJsonPayload(rawText):
    cleaned = rawText.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.split("\n", 1)[-1]
        cleaned = cleaned.rsplit("```", 1)[0]
    return json.loads(cleaned.strip())


def analyzeEvent(apiKey, description):
    client = OpenAI(
        api_key=apiKey,
        base_url=DEEPSEEK_BASE_URL,
        timeout=API_TIMEOUT_SECONDS,
    )
    response = client.chat.completions.create(
        model=DEEPSEEK_MODEL,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {
                "role": "user",
                "content": f"请对以下护理不良事件描述进行根本原因分析：\n\n{description.strip()}",
            },
        ],
        response_format={"type": "json_object"},
        temperature=0.2,
    )
    return parseJsonPayload(response.choices[0].message.content)


def buildRecommendationFrame(recommendations):
    rows = []
    for item in recommendations if isinstance(recommendations, list) else []:
        if isinstance(item, dict):
            measure = item.get("措施") or item.get("measure") or str(item)
            priority = item.get("优先级") or item.get("priority") or "中"
        else:
            measure, priority = str(item), "中"
        rows.append({"改进措施": str(measure), "优先级": str(priority)})

    return pd.DataFrame(rows, columns=["改进措施", "优先级"])


def renderAnalysisResult(result, description):
    if not result.get("is_event", False):
        st.info(
            result.get("message")
            or "该内容看起来不是护理不良事件描述，本系统仅用于护理不良事件的根本原因分析。请粘贴包含事件经过、涉及环节的临床记录后重试。"
        )
        return

    severity = str(result.get("severity", "")).strip()
    typeColumn, severityColumn = st.columns(2)
    with typeColumn:
        st.markdown("**事件类型**")
        st.markdown(f"### {result.get('event_type') or '未识别'}")
    with severityColumn:
        st.markdown("**严重程度**")
        st.markdown(f"### {SEVERITY_BADGES.get(severity, severity or '未评估')}")

    deepAnalysis = str(result.get("deep_analysis", "")).strip()
    if deepAnalysis:
        st.markdown("**深度原因剖析**")
        with st.container(border=True):
            st.markdown(deepAnalysis)

    st.markdown("**根本原因（4M1E）**")
    rootCauses = result.get("root_causes")
    if isinstance(rootCauses, dict) and rootCauses:
        for dimension, causes in rootCauses.items():
            with st.container(border=True):
                st.markdown(f"**{dimension}**")
                if isinstance(causes, list) and causes:
                    for cause in causes:
                        st.markdown(f"- {cause}")
                else:
                    st.markdown("- 未识别到该类原因")
    else:
        st.caption("未识别到明确的根本原因。")

    st.markdown("**改进措施建议**")
    recommendationFrame = buildRecommendationFrame(result.get("recommendations"))
    if recommendationFrame.empty:
        st.caption("未生成改进措施建议。")
    else:
        st.dataframe(
            recommendationFrame,
            width="stretch",
            hide_index=True,
            column_config={
                "改进措施": st.column_config.TextColumn(width="large"),
                "优先级": st.column_config.TextColumn(width="small"),
            },
        )

    with st.expander("查看原始 JSON 结果"):
        st.json(result)

    renderFeedbackModule(result, description)


def resultFingerprint(result):
    payload = json.dumps(result, ensure_ascii=False, sort_keys=True)
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()[:10]


def splitNonEmptyLines(text):
    return [line.strip() for line in text.splitlines() if line.strip()]


def feedbackCount():
    if not FEEDBACK_PATH.is_file():
        return 0
    try:
        records = json.loads(FEEDBACK_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return 0
    return len(records) if isinstance(records, list) else 0


def appendFeedbackRecord(record):
    records = []
    if FEEDBACK_PATH.is_file():
        raw = FEEDBACK_PATH.read_text(encoding="utf-8")
        if raw.strip():
            records = json.loads(raw)
            if not isinstance(records, list):
                raise ValueError("user_feedback.json 顶层结构不是 JSON 数组")

    updated = records + [record]
    tempPath = FEEDBACK_PATH.with_name(f"{FEEDBACK_PATH.name}.tmp")
    tempPath.write_text(json.dumps(updated, ensure_ascii=False, indent=2), encoding="utf-8")
    tempPath.replace(FEEDBACK_PATH)
    return len(updated)


def clearFeedbackEditorState():
    for key in [key for key in st.session_state if key.startswith(EDITOR_KEY_PREFIX)]:
        del st.session_state[key]


def seedEditorState(key, defaultValue):
    if key not in st.session_state:
        st.session_state[key] = defaultValue
    return key


def renderFeedbackModule(result, description):
    st.divider()
    st.subheader("✏️ 修改AI分析结果")
    st.caption("AI 输出仅作参考。可在此逐条修正根本原因与改进措施，保存后会写入本地 user_feedback.json，用于人工复核与后续优化。")

    fingerprint = resultFingerprint(result)
    rootCauses = result.get("root_causes") if isinstance(result.get("root_causes"), dict) else {}
    dimensions = [name for name in RCA_DIMENSIONS if name in rootCauses] or list(rootCauses) or RCA_DIMENSIONS

    editedCauses = {}
    with st.expander("编辑根本原因（每行一条）", expanded=True):
        for dimension in dimensions:
            causes = rootCauses.get(dimension)
            defaultText = "\n".join(str(item) for item in causes) if isinstance(causes, list) else ""
            editedCauses[dimension] = st.text_area(
                dimension,
                key=seedEditorState(f"{EDITOR_KEY_PREFIX}cause_{fingerprint}_{dimension}", defaultText),
                height=100,
            )

    recommendationItems = result.get("recommendations")
    recommendationLines = (
        [
            str(item.get("措施")) if isinstance(item, dict) and item.get("措施") else str(item)
            for item in recommendationItems
        ]
        if isinstance(recommendationItems, list)
        else []
    )
    editedRecommendations = st.text_area(
        "改进措施（每行一条）",
        key=seedEditorState(
            f"{EDITOR_KEY_PREFIX}recommendations_{fingerprint}",
            "\n".join(recommendationLines),
        ),
        height=140,
    )
    editorNote = st.text_area(
        "修改说明（可选）",
        key=seedEditorState(f"{EDITOR_KEY_PREFIX}note_{fingerprint}", ""),
        height=80,
        placeholder="例如：结合科室实际情况删去不具备可行性的措施",
    )

    if st.button("💾 保存修正结果"):
        record = {
            "savedAt": datetime.now().isoformat(timespec="seconds"),
            "inputDescription": description.strip(),
            "originalEventType": result.get("event_type", ""),
            "originalSeverity": result.get("severity", ""),
            "originalRootCauses": rootCauses,
            "originalRecommendations": recommendationLines,
            "correctedRootCauses": {
                dimension: splitNonEmptyLines(text) for dimension, text in editedCauses.items()
            },
            "correctedRecommendations": splitNonEmptyLines(editedRecommendations),
            "editorNote": editorNote.strip(),
        }

        try:
            total = appendFeedbackRecord(record)
            st.success(f"修正结果已保存，累计 {total} 条人工反馈。")
        except json.JSONDecodeError:
            st.error("user_feedback.json 内容已损坏，无法解析。请先备份并删除该文件后重新保存。")
        except (OSError, ValueError) as error:
            st.error(f"保存失败：{error}")


def renderAnalysisPage(apiKey):
    st.header("智能根因分析")

    with st.container(border=True):
        st.markdown(f"⚠️ **免责声明**：{DISCLAIMER}")

    st.info("粘贴一段不良事件文字描述，系统将基于「人、机、料、法、环（4M1E）」框架自动生成根因分析与改进措施。")

    description = st.text_area(
        "不良事件描述",
        height=180,
        placeholder="例如：夜班护士在给3床患者发药时，因同时接听电话导致发错药，患者服用后出现轻微不良反应……",
    )

    if st.button("开始智能根因分析", type="primary"):
        if not apiKey:
            st.warning("请先配置 API Key")
        elif not description.strip():
            st.warning("请先粘贴不良事件的文字描述")
        else:
            try:
                with st.spinner("AI正在深度思考中..."):
                    clearFeedbackEditorState()
                    st.session_state["analysisResult"] = analyzeEvent(apiKey, description)
                    st.session_state["analysisInput"] = description.strip()
            except AuthenticationError:
                st.error("API Key 无效或已失效，请检查 config.json 中的密钥后重试。")
            except RateLimitError:
                st.error("请求过于频繁或账户额度不足，请稍后重试。")
            except APITimeoutError:
                st.error("请求超时，请检查网络连接后重试。")
            except APIConnectionError:
                st.error("无法连接 DeepSeek 服务，请检查网络或代理设置。")
            except APIStatusError as error:
                st.error(f"DeepSeek 接口返回错误（状态码 {error.status_code}）：{error.message}")
            except (json.JSONDecodeError, ValueError) as error:
                st.error(f"AI 返回内容解析失败，请重试。详细信息：{error}")

    if "analysisResult" in st.session_state:
        st.divider()
        renderAnalysisResult(
            st.session_state["analysisResult"],
            st.session_state.get("analysisInput", ""),
        )


def textColorFor(hexColor):
    red, green, blue = (
        int(hexColor[index:index + 2], 16) / 255 for index in (1, 3, 5)
    )
    luminance = 0.2126 * red + 0.7152 * green + 0.0722 * blue
    return "#0B0B0B" if luminance > 0.55 else "#FFFFFF"


def applyChartStyle(figure, title, style):
    figure.update_layout(
        title=dict(text=title, x=0, xanchor="left", font=dict(size=14, color=style["ink"])),
        paper_bgcolor=style["surface"],
        plot_bgcolor=style["surface"],
        font=dict(family=FONT_FAMILY, color=style["muted"], size=12),
        margin=dict(l=8, r=8, t=52, b=8),
        height=CHART_HEIGHT,
        showlegend=False,
        bargap=0.35,
        hoverlabel=dict(
            bgcolor=style["surface"],
            bordercolor=style["axis"],
            font=dict(family=FONT_FAMILY, color=style["ink"], size=12),
        ),
    )
    figure.update_xaxes(
        showgrid=False,
        zeroline=False,
        linecolor=style["axis"],
        tickfont=dict(color=style["muted"]),
    )
    figure.update_yaxes(
        showgrid=True,
        gridcolor=style["grid"],
        gridwidth=1,
        zeroline=False,
        linecolor=style["axis"],
        tickfont=dict(color=style["muted"]),
    )
    return figure


def buildDepartmentChart(filteredFrame, style):
    counts = filteredFrame["科室"].dropna().value_counts().sort_values(ascending=False)
    if counts.empty:
        return None

    figure = go.Figure(
        go.Bar(
            x=counts.index.tolist(),
            y=counts.values.tolist(),
            marker=dict(color=style["palette"][0]),
            hovertemplate="%{x}<br>事件数 %{y} 起<extra></extra>",
        )
    )
    figure.update_xaxes(tickangle=-40)
    return DEPARTMENT_CHART_TITLE, applyChartStyle(figure, DEPARTMENT_CHART_TITLE, style)


def buildTypeColorMap(eventFrame, style):
    topTypes = eventFrame["事件类型"].value_counts().index[:MAX_PIE_SLICES].tolist()
    return {name: style["palette"][index] for index, name in enumerate(topTypes)}


def buildTypeChart(filteredFrame, typeColorMap, style):
    labels = filteredFrame["事件类型"].dropna().map(
        lambda name: name if name in typeColorMap else OTHER_LABEL
    )
    counts = labels.value_counts()
    if counts.empty:
        return None

    colors = [typeColorMap.get(label, style["other"]) for label in counts.index]

    figure = go.Figure(
        go.Pie(
            labels=counts.index.tolist(),
            values=counts.values.tolist(),
            hole=0.45,
            sort=True,
            marker=dict(colors=colors, line=dict(color=style["surface"], width=2)),
            textinfo="label+percent",
            textposition="auto",
            insidetextfont=dict(
                family=FONT_FAMILY,
                color=[textColorFor(color) for color in colors],
                size=12,
            ),
            outsidetextfont=dict(family=FONT_FAMILY, color=style["muted"], size=12),
            hovertemplate="%{label}<br>事件数 %{value} 起<br>占比 %{percent}<extra></extra>",
        )
    )
    return TYPE_CHART_TITLE, applyChartStyle(figure, TYPE_CHART_TITLE, style)


def buildSeverityChart(filteredFrame, style):
    severityCounts = (
        filteredFrame["严重程度"].dropna().value_counts()
        if "严重程度" in filteredFrame.columns
        else pd.Series(dtype="int64")
    )

    if not severityCounts.empty:
        levels = [level for level in SEVERITY_LEVELS if level in severityCounts.index]
        levels += [level for level in severityCounts.index if level not in SEVERITY_LEVELS]
        values = [int(severityCounts[level]) for level in levels]
        colors = [
            style["ordinal"][SEVERITY_LEVELS.index(level)]
            if level in SEVERITY_LEVELS
            else style["other"]
            for level in levels
        ]
        title = SEVERITY_CHART_TITLE
    else:
        counts = filteredFrame["事件类型"].dropna().value_counts()
        if counts.empty:
            return None
        levels = counts.index.tolist()
        values = counts.values.tolist()
        colors = [style["palette"][0]] * len(levels)
        title = FALLBACK_TYPE_CHART_TITLE

    figure = go.Figure(
        go.Bar(
            x=values[::-1],
            y=levels[::-1],
            orientation="h",
            marker=dict(color=colors[::-1]),
            text=[str(value) for value in values[::-1]],
            textposition="outside",
            textfont=dict(family=FONT_FAMILY, color=style["muted"], size=12),
            hovertemplate="%{y}<br>事件数 %{x} 起<extra></extra>",
        )
    )
    figure.update_xaxes(range=[0, max(values) * 1.25])
    return title, applyChartStyle(figure, title, style)


def buildChartSpecs(filteredFrame, eventFrame, style):
    typeColorMap = buildTypeColorMap(eventFrame, style)
    results = [
        buildDepartmentChart(filteredFrame, style),
        buildTypeChart(filteredFrame, typeColorMap, style),
        buildSeverityChart(filteredFrame, style),
    ]
    return [result for result in results if result is not None]


def formatShare(count, total):
    return f"{count / total:.1%}" if total else "0.0%"


def buildOverview(filteredFrame, rangeLabel):
    severityCounts = (
        filteredFrame["严重程度"].dropna().value_counts()
        if "严重程度" in filteredFrame.columns
        else None
    )

    overview = {
        "total": len(filteredFrame),
        "departmentCount": int(filteredFrame["科室"].dropna().nunique()),
        "topDepartments": [
            (name, int(count))
            for name, count in filteredFrame["科室"].dropna().value_counts().head(5).items()
        ],
        "topTypes": [
            (name, int(count))
            for name, count in filteredFrame["事件类型"].dropna().value_counts().head(3).items()
        ],
        "severityCounts": [],
        "severeCount": None,
        "moderatePlusCount": None,
        "rangeLabel": rangeLabel,
        "dataSourceLabel": formatDataSourceLabel(),
    }

    if severityCounts is not None and not severityCounts.empty:
        overview["severeCount"] = int(severityCounts.get("重度", 0))
        overview["moderatePlusCount"] = int(
            severityCounts.reindex(["中度", "重度"]).fillna(0).sum()
        )
        overview["severityCounts"] = [
            (level, int(severityCounts[level]))
            for level in SEVERITY_LEVELS
            if level in severityCounts.index
        ]

    return overview


def buildFindingsText(overview):
    total = overview["total"]
    sentences = [
        f"本期共纳入分析护理不良事件 {total} 起，覆盖 {overview['departmentCount']} 个科室。"
    ]

    if overview["topDepartments"]:
        name, count = overview["topDepartments"][0]
        sentences.append(
            f"其中「{name}」上报 {count} 起，占 {formatShare(count, total)}，居各科室首位，建议列为下阶段质控重点监测对象。"
        )

    if overview["topTypes"]:
        typeDetail = "、".join(
            f"「{name}」{count} 起（{formatShare(count, total)}）"
            for name, count in overview["topTypes"]
        )
        coveredShare = sum(count for _, count in overview["topTypes"])
        sentences.append(
            f"事件类型以「{overview['topTypes'][0][0]}」为主，前三类为{typeDetail}，合计占 {formatShare(coveredShare, total)}，是本轮干预的核心方向。"
        )

    if overview["moderatePlusCount"] is not None:
        sentences.append(
            f"从严重程度看，中度及以上事件 {overview['moderatePlusCount']} 起，占 {formatShare(overview['moderatePlusCount'], total)}，其中重度 {overview['severeCount']} 起，需优先落实整改并跟踪闭环。"
        )
    else:
        sentences.append(
            "当前数据未包含「严重程度」字段，建议在不良事件上报表单中补充严重程度分级，以便开展风险分层管理。"
        )

    sentences.append(
        "建议围绕高发科室、高发类型的共性环节（查对制度执行、夜班人力配置、管路固定与约束具使用、巡视与交接班落实）制定专项改进计划，并在下一个质控周期复评同类事件发生率。"
    )
    return "".join(sentences)


def pickTypicalCase(filteredFrame):
    if filteredFrame.empty:
        return None

    if "严重程度" in filteredFrame.columns:
        severityRank = filteredFrame["严重程度"].map(
            {level: index for index, level in enumerate(SEVERITY_LEVELS)}
        )
        ranked = severityRank.dropna()
        if not ranked.empty:
            return filteredFrame.loc[ranked.sort_values(ascending=False).index[0]]

    days = pd.to_datetime(filteredFrame["发生日期"], errors="coerce")
    if days.notna().any():
        return filteredFrame.loc[days.sort_values(ascending=False).index[0]]
    return filteredFrame.iloc[0]


def generateRectificationAdvice(apiKey, caseRecord):
    if not apiKey:
        return None, "未检测到可用的 API Key，本段 AI 整改建议未生成。请在侧边栏配置密钥后重新生成报告。"

    caseText = "\n".join(
        [
            f"发生科室：{caseRecord.get('科室', '未填写')}",
            f"事件类型：{caseRecord.get('事件类型', '未填写')}",
            f"严重程度：{caseRecord.get('严重程度', '未评估')}",
            f"事件经过：{caseRecord.get('事件经过描述', '未填写')}",
        ]
    )

    try:
        client = OpenAI(
            api_key=apiKey,
            base_url=DEEPSEEK_BASE_URL,
            timeout=API_TIMEOUT_SECONDS,
        )
        response = client.chat.completions.create(
            model=DEEPSEEK_MODEL,
            messages=[
                {"role": "system", "content": RECTIFICATION_PROMPT},
                {"role": "user", "content": f"请针对以下护理不良事件给出整改建议：\n\n{caseText}"},
            ],
            temperature=0.3,
        )
        return response.choices[0].message.content.strip(), None
    except AuthenticationError:
        return None, "API Key 无效或已失效，AI 整改建议生成失败，请在侧边栏检查密钥后重新生成报告。"
    except RateLimitError:
        return None, "请求过于频繁或账户额度不足，AI 整改建议生成失败，请稍后重新生成报告。"
    except APITimeoutError:
        return None, "请求超时，AI 整改建议生成失败，请检查网络后重新生成报告。"
    except APIConnectionError:
        return None, "无法连接 DeepSeek 服务，AI 整改建议生成失败，请检查网络或代理设置后重新生成报告。"
    except APIStatusError as error:
        return None, f"DeepSeek 接口返回错误（状态码 {error.status_code}），AI 整改建议生成失败，请稍后重新生成报告。"
    except (ValueError, OSError) as error:
        return None, f"AI 整改建议生成失败：{error}"


def exportChartImage(figure):
    try:
        return figure.to_image(
            format="png",
            width=CHART_IMAGE_WIDTH_PX,
            height=CHART_IMAGE_HEIGHT_PX,
            scale=CHART_IMAGE_SCALE,
        )
    except (ValueError, RuntimeError, ImportError, OSError):
        return None


def configureDocumentStyles(document):
    bodyStyle = document.styles["Normal"]
    bodyStyle.font.name = DOC_LATIN_FONT
    bodyStyle.font.size = Pt(10.5)
    bodyStyle.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOC_BODY_FONT)

    titleStyle = document.styles["Title"]
    titleStyle.font.name = DOC_LATIN_FONT
    titleStyle.font.size = Pt(22)
    titleStyle.font.color.rgb = DOC_ACCENT_COLOR
    titleStyle.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), DOC_HEADING_FONT
    )

    for styleName, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        headingStyle = document.styles[styleName]
        headingStyle.font.name = DOC_LATIN_FONT
        headingStyle.font.size = Pt(size)
        headingStyle.font.color.rgb = DOC_ACCENT_COLOR
        headingStyle.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), DOC_HEADING_FONT
        )


def addMetricParagraph(document, label, value):
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}：").bold = True
    paragraph.add_run(str(value))
    return paragraph


def addCaptionParagraph(document, text):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = DOC_MUTED_COLOR
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def addTableRows(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    headerCells = table.rows[0].cells
    for cell, header in zip(headerCells, headers):
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)

    document.add_paragraph()
    return table


def normalizeAdviceLine(line):
    cleaned = re.sub(r"^[#>\-\*•\s]+", "", line.strip())
    return cleaned.replace("**", "").strip()


def buildReportDocument(overview, findings, caseRecord, advice, adviceWarning, chartSpecs):
    document = Document()
    configureDocumentStyles(document)

    document.add_heading(REPORT_TITLE, level=0)
    addCaptionParagraph(
        document,
        f"数据来源：{overview['dataSourceLabel']}　|　统计区间：{overview['rangeLabel']}　|　生成时间：{datetime.now():%Y-%m-%d %H:%M}",
    )

    document.add_heading("一、数据概览", level=1)
    addMetricParagraph(document, "事件总数", f"{overview['total']} 起")
    addMetricParagraph(document, "涉及科室", f"{overview['departmentCount']} 个")
    if overview["topDepartments"]:
        name, count = overview["topDepartments"][0]
        addMetricParagraph(
            document,
            "事件最多的科室",
            f"{name}（{count} 起，占 {formatShare(count, overview['total'])}）",
        )
    if overview["topTypes"]:
        addMetricParagraph(
            document,
            "主要事件类型",
            "、".join(f"{name}（{count} 起）" for name, count in overview["topTypes"]),
        )
    if overview["moderatePlusCount"] is not None:
        addMetricParagraph(
            document,
            "中度及以上事件",
            f"{overview['moderatePlusCount']} 起（占 {formatShare(overview['moderatePlusCount'], overview['total'])}），其中重度 {overview['severeCount']} 起",
        )
    else:
        addMetricParagraph(document, "严重程度分级", "当前数据未包含该字段")
    addMetricParagraph(document, "统计区间", overview["rangeLabel"])

    if overview["severityCounts"]:
        document.add_heading("严重程度构成", level=2)
        addTableRows(
            document,
            ["严重程度", "事件数（起）", "占比"],
            [
                [level, count, formatShare(count, overview["total"])]
                for level, count in overview["severityCounts"]
            ],
        )

    if overview["topDepartments"]:
        document.add_heading("科室分布（前五位）", level=2)
        addTableRows(
            document,
            ["科室", "事件数（起）", "占比"],
            [
                [name, count, formatShare(count, overview["total"])]
                for name, count in overview["topDepartments"]
            ],
        )

    document.add_heading("二、主要发现", level=1)
    document.add_paragraph(findings)

    document.add_heading("三、典型案例与整改建议", level=1)
    if caseRecord is None:
        document.add_paragraph("当前筛选条件下没有可用于分析的典型案例。")
    else:
        document.add_heading("事件基本情况", level=2)
        addMetricParagraph(document, "事件编号", caseRecord.get("事件ID", "未填写"))
        occurredOn = pd.to_datetime(caseRecord.get("发生日期"), errors="coerce")
        addMetricParagraph(
            document,
            "发生日期",
            occurredOn.strftime("%Y-%m-%d") if pd.notna(occurredOn) else "未填写",
        )
        addMetricParagraph(document, "发生科室", caseRecord.get("科室", "未填写"))
        addMetricParagraph(document, "事件类型", caseRecord.get("事件类型", "未填写"))
        addMetricParagraph(document, "严重程度", caseRecord.get("严重程度", "未评估"))
        document.add_paragraph(str(caseRecord.get("事件经过描述", "未填写")))

        document.add_heading("整改建议", level=2)
        if advice:
            for line in advice.splitlines():
                cleanedLine = normalizeAdviceLine(line)
                if cleanedLine:
                    document.add_paragraph(cleanedLine)
        else:
            document.add_paragraph(adviceWarning or "未生成整改建议。")

    document.add_heading("四、可视化图表", level=1)
    imageCount = 0
    for title, figure in chartSpecs:
        imageBytes = exportChartImage(figure)
        if imageBytes is None:
            continue
        document.add_picture(io.BytesIO(imageBytes), width=Inches(DOC_IMAGE_WIDTH_INCHES))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        addCaptionParagraph(document, title)
        imageCount += 1

    if imageCount == 0:
        document.add_paragraph(
            "本次未能导出图表图片（未安装 kaleido 或导出失败），报告仅输出文字内容。"
        )

    document.add_heading("五、免责声明", level=1)
    document.add_paragraph(DISCLAIMER)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def renderReportExport(filteredFrame, eventFrame, overview, apiKey):
    st.divider()
    st.subheader("报告导出")
    st.caption("报告包含数据概览、主要发现、典型案例整改建议与可视化图表，可直接用于护理质控会议。")

    if st.button("📄 一键生成Word分析报告", type="primary"):
        try:
            with st.spinner("正在生成报告，请稍候..."):
                caseRecord = pickTypicalCase(filteredFrame)
                advice, adviceWarning = None, "当前筛选条件下没有可用于分析的典型案例。"
                if caseRecord is not None:
                    advice, adviceWarning = generateRectificationAdvice(apiKey, caseRecord)
                reportBytes = buildReportDocument(
                    overview,
                    buildFindingsText(overview),
                    caseRecord,
                    advice,
                    adviceWarning,
                    buildChartSpecs(filteredFrame, eventFrame, DOCUMENT_STYLE),
                )

            st.session_state["reportBytes"] = reportBytes
            st.session_state["reportFileName"] = (
                f"NurseInsight_护理不良事件分析报告_{datetime.now():%Y%m%d_%H%M}.docx"
            )
            st.success("报告已生成，可点击下方按钮下载。")
        except (OSError, ValueError) as error:
            st.error(f"报告生成失败：{error}")

    if "reportBytes" in st.session_state:
        st.download_button(
            "⬇️ 下载 Word 报告（.docx）",
            data=st.session_state["reportBytes"],
            file_name=st.session_state["reportFileName"],
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        st.caption("报告内容对应生成时的筛选条件，若已调整筛选请重新生成。")


def renderDashboardPage(apiKey):
    st.header("可视化质控看板")

    eventFrame = getEventData()
    if eventFrame.empty:
        st.warning("当前没有可分析的数据，请先在「数据管理与录入」页面上传数据。")
        return

    st.caption(f"数据来源：{formatDataSourceLabel()}")

    allDepartments = sorted(eventFrame["科室"].dropna().unique().tolist())
    allDays = pd.to_datetime(eventFrame["发生日期"], errors="coerce").dropna()

    filterColumns = st.columns([2, 2])
    with filterColumns[0]:
        selectedDepartments = st.multiselect("科室筛选", allDepartments, default=allDepartments)
    with filterColumns[1]:
        if allDays.empty:
            selectedRange = None
            st.caption("当前数据缺少有效日期，已跳过时间筛选")
        else:
            minDay, maxDay = allDays.min().date(), allDays.max().date()
            pickedRange = st.date_input(
                "发生日期区间",
                value=(minDay, maxDay),
                min_value=minDay,
                max_value=maxDay,
            )
            selectedRange = pickedRange if len(pickedRange) == 2 else (minDay, maxDay)

    filteredFrame = eventFrame
    if selectedDepartments and len(selectedDepartments) < len(allDepartments):
        filteredFrame = filteredFrame[filteredFrame["科室"].isin(selectedDepartments)]
    if selectedRange is not None:
        occurredDays = pd.to_datetime(filteredFrame["发生日期"], errors="coerce")
        startDay, endDay = pd.Timestamp(selectedRange[0]), pd.Timestamp(selectedRange[1])
        filteredFrame = filteredFrame[(occurredDays >= startDay) & (occurredDays <= endDay)]
        rangeLabel = f"{startDay:%Y-%m-%d} 至 {endDay:%Y-%m-%d}"
    else:
        rangeLabel = "全部数据（无有效日期）"

    if filteredFrame.empty:
        st.warning("当前筛选条件下没有数据，请放宽筛选范围。")
        return

    topDepartment = filteredFrame["科室"].value_counts()
    severeCount = (
        int((filteredFrame["严重程度"] == "重度").sum())
        if "严重程度" in filteredFrame.columns
        else 0
    )

    metricColumns = st.columns(4)
    metricColumns[0].metric("事件总数", f"{len(filteredFrame)} 起")
    metricColumns[1].metric("涉及科室", f"{filteredFrame['科室'].nunique()} 个")
    metricColumns[2].metric("事件最多的科室", topDepartment.index[0], f"{topDepartment.iloc[0]} 起")
    metricColumns[3].metric(
        "重度事件",
        f"{severeCount} 起" if "严重程度" in filteredFrame.columns else "无该字段",
        f"{severeCount / len(filteredFrame):.1%}" if "严重程度" in filteredFrame.columns else None,
    )

    st.write("")

    chartSpecs = buildChartSpecs(filteredFrame, eventFrame, SCREEN_STYLE)

    for column, spec in zip(st.columns(3, gap="large"), chartSpecs):
        with column:
            st.plotly_chart(
                spec[1],
                width="stretch",
                config={"displayModeBar": False},
            )

    with st.expander("查看数据表（表格视图）"):
        tableColumns = [
            column
            for column in ["事件ID", "发生日期", "科室", "事件类型", "严重程度"]
            if column in filteredFrame.columns
        ]
        st.dataframe(
            filteredFrame[tableColumns].sort_values("发生日期", ascending=False),
            width="stretch",
            hide_index=True,
        )

    renderReportExport(
        filteredFrame,
        eventFrame,
        buildOverview(filteredFrame, rangeLabel),
        apiKey,
    )


def renderPage(pageName, apiKey):
    if pageName == DATA_PAGE:
        renderDataPage()
    elif pageName == ANALYSIS_PAGE:
        renderAnalysisPage(apiKey)
    else:
        renderDashboardPage(apiKey)


def main():
    st.set_page_config(page_title=APP_TITLE, page_icon="🏥", layout="wide")
    st.markdown(GLOBAL_CSS, unsafe_allow_html=True)
    st.title(APP_TITLE)

    apiKey, keySource = loadConfiguredApiKey()

    with st.sidebar:
        st.markdown("### 导航")
        selectedPage = st.radio("功能模块", PAGES, label_visibility="collapsed")

        st.divider()
        if apiKey:
            st.caption(f"✅ 已启用密钥（来源：{keySource}），密钥内容不显示")
        else:
            st.caption("未检测到本地配置，请在下方手动输入 API Key")
            apiKey = st.text_input(
                "DeepSeek API Key",
                type="password",
                help="未检测到本地 config.json 或 Streamlit Secrets，请在此手动输入密钥，仅保存在当前会话中。",
            )

        savedFeedback = feedbackCount()
        if savedFeedback:
            st.caption(f"📝 已收集人工修正 {savedFeedback} 条")

        if str(st.get_option("theme.base")).lower() != "dark":
            st.caption(MISSING_THEME_HINT)

        with st.expander("🔐 打包与分享前必读"):
            st.markdown(
                "- 密钥来源为 `config.json` 或 Streamlit Secrets，界面**不会显示密钥内容**。\n"
                "- **如果要打包发给别人，请务必删除 `config.json`，或将其重命名为 `config.example.json` 并清空密钥。**\n"
                "- 部署到 Streamlit Cloud 时，请在 Secrets 面板配置 `DEEPSEEK_API_KEY`，不要提交 `config.json`。\n"
                "- `config.json`、`user_feedback.json` 已列入 `.gitignore`。"
            )

        st.divider()
        st.caption(DISCLAIMER)

    renderPage(selectedPage, apiKey)


if __name__ == "__main__":
    main()
